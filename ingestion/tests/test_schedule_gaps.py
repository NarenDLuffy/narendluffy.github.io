"""Regression: Wednesday 14 Oct 2026, Tamna BC Hall 17:00-19:30 (RAN1#126bis).

The official grid has a Tuesday "Early dinner" sub-row under 17:00 ~ 19:30.
Its other cells are vertical-merge continuations that Word stores without the
lane colour, so they were read a second time as white (Main) cells: Yeongju A's
"Modulation" landed in Tamna BC and the merge dropped 10.4.2 and 10.5.5.
"""

from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ingestion.block_schedule import parse_block_schedule_docx
from ingestion.live import coverage_gaps
from ingestion.models import Room, ScheduleSource, Session, SessionSourceRef

SRC = ScheduleSource(
    sourceId="s", meetingId="m", fileName="f.docx", label="f", type="x",
    origin="public", retrievedAt="x",
)
FILLS = ["FFFFFF", "D9D9D9"]


def _cell_props(cell, fill: str | None, merge: str | None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)
    if merge:
        v = OxmlElement("w:vMerge")
        if merge == "restart":
            v.set(qn("w:val"), "restart")
        tc_pr.append(v)


def _grid(path: str) -> None:
    doc = Document()
    doc.add_paragraph("RAN1#126b Online Session Schedule")
    table = doc.add_table(rows=3, cols=3)
    head = table.rows[0].cells
    head[0].text, head[1].text = "", "Wednesday"
    head[1].merge(head[2])
    row = table.rows[1].cells
    row[0].text = "17:00 ~ 19:30 (150 min)"
    row[1].text = "6GR (150)\n.10.4.1 (40)\n.10.4.2 (40)\n.10.5.5 (40)\n.10.5.3.1 (30)"
    row[2].text = "6GR (150)\n.10.3.1 Channel coding (60)\n.10.3.2 Modulation (30)\n.10.5.2.2 (60)"
    _cell_props(row[0], "FFFFFF", "restart")
    _cell_props(row[1], FILLS[0], "restart")
    _cell_props(row[2], FILLS[1], "restart")
    # Continuation sub-row: every cell continues the one above; Word stores
    # them white, losing the lane colour (as in the real file).
    for cell in table.rows[2].cells:
        _cell_props(cell, "FFFFFF", "continue")
    doc.save(path)


def test_continuation_rows_are_not_read_twice(tmp_path):
    path = str(tmp_path / "grid.docx")
    _grid(path)
    rooms, sessions = parse_block_schedule_docx(
        path, meeting_id="m", start_date="2026-10-12", end_date="2026-10-16", source=SRC
    )
    white = next(r for r in rooms if r.sourceColor == "FFFFFF")
    main = sorted(
        (s.startTime, s.endTime, s.agendaItems[0])
        for s in sessions
        if s.roomId == white.roomId and s.date == "2026-10-14" and s.kind == "session"
    )
    assert main == [
        ("17:00", "17:40", "10.4.1"),
        ("17:40", "18:20", "10.4.2"),
        ("18:20", "19:00", "10.5.5"),
        ("19:00", "19:30", "10.5.3.1"),
    ]


def _s(start: str, end: str, room: str = "r", code: str = "10.4.1") -> Session:
    return Session(
        sessionId=f"{room}{start}", meetingId="m", date="2026-10-14", day="Wednesday",
        startTime=start, endTime=end, roomId=room, roomName="Tamna BC Hall", topic=code,
        topicKey=code, agendaItems=[code], kind="session", status="scheduled",
        sources=[SessionSourceRef("s", ["time"])],
    )


def test_a_dropped_item_is_reported_as_a_gap():
    rooms = [Room(roomId="r", meetingId="m", roomName="Tamna BC Hall", order=0)]
    parsed = [_s("17:00", "17:40"), _s("17:40", "18:20"), _s("18:20", "19:00"), _s("19:00", "19:30")]
    published = [_s("17:00", "17:40"), _s("19:00", "19:30")]
    gaps = coverage_gaps(parsed, published, rooms)
    assert gaps == [
        "2026-10-14 Tamna BC Hall: no session 17:40-19:00 although a schedule document fills it"
    ]
    assert coverage_gaps(parsed, parsed, rooms) == []


def test_one_cell_read_into_two_rooms_is_refused():
    from dataclasses import replace

    from ingestion.live import cell_checks

    rooms = [
        Room(roomId="r", meetingId="m", roomName="Tamna BC Hall", order=0),
        Room(roomId="y", meetingId="m", roomName="Yeongju A Hall", order=1),
    ]
    good = _s("18:00", "18:30", room="y", code="10.3.2")
    good.sources = [SessionSourceRef("s", ["time"], cell="s|t3|Wednesday|1020|c2|ab")]
    copy = replace(good, roomId="r", sources=[SessionSourceRef("s", ["time"], cell=good.sources[0].cell)])
    errors = cell_checks([good, copy], [good, copy], rooms)
    assert any("one table cell was read into 2 rooms" in e for e in errors)
    assert any("10.3.2 is in both" in e for e in errors)
    assert cell_checks([good], [good], rooms) == []


def test_item_listed_in_two_rooms_by_the_document_is_allowed():
    from ingestion.live import cell_checks

    rooms = [
        Room(roomId="r", meetingId="m", roomName="Tamna BC Hall", order=0),
        Room(roomId="y", meetingId="m", roomName="Yeongju A Hall", order=1),
    ]
    a = _s("18:00", "18:30", room="r", code="10.3.2")
    a.sources = [SessionSourceRef("s", ["time"], cell="s|t3|Wednesday|1020|c1|aa")]
    b = _s("18:00", "18:30", room="y", code="10.3.2")
    b.sources = [SessionSourceRef("s", ["time"], cell="s|t3|Wednesday|1020|c2|bb")]
    assert cell_checks([a, b], [a, b], rooms) == []
